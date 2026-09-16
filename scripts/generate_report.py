import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def build_rows(detail):
    checks = detail.get("checks", [])
    rows = [["Field", "Status", "Rule", "Source", "Confidence"]]
    for check in checks:
        source = check.get("source") or "Source support unavailable"
        if check.get("page"):
            source = f"{source}, page {check['page']}"
        rows.append([
            clean(check.get("label") or check.get("field")),
            clean(check.get("status")),
            clean(check.get("rule")),
            source,
            f"{round(float(check.get('confidence') or 0) * 100)}%",
        ])
    return rows


def create_pdf(detail, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = []
    inspection = detail.get("inspection", {})
    declaration = (detail.get("declaration") or {}).get("declarations") or {}

    story.append(Paragraph("PackCheck AI Inspection Report", styles["Title"]))
    story.append(Paragraph(f"Inspection ID: {inspection.get('id', '')}", styles["Normal"]))
    story.append(Paragraph(f"Status: {inspection.get('status', '')}", styles["Normal"]))
    story.append(Paragraph(f"Result: {inspection.get('result_status', '')}", styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Extracted Declarations", styles["Heading2"]))
    declaration_rows = [["Field", "Value"]] + [[key.replace("_", " "), clean(value) or "Not Detected"] for key, value in declaration.items()]
    table = Table(declaration_rows, colWidths=[170, 310])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
    ]))
    story.append(table)
    story.append(Spacer(1, 14))

    story.append(Paragraph("Compliance Checks", styles["Heading2"]))
    checks_table = Table(build_rows(detail), colWidths=[95, 80, 135, 125, 45])
    checks_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEADING", (0, 0), (-1, -1), 9),
    ]))
    story.append(checks_table)
    doc.build(story)


def create_docx(detail, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    inspection = detail.get("inspection", {})
    declaration = (detail.get("declaration") or {}).get("declarations") or {}

    title = doc.add_paragraph()
    run = title.add_run("PackCheck AI Inspection Report")
    run.bold = True
    run.font.size = docx_pt(20)
    doc.add_paragraph(f"Inspection ID: {inspection.get('id', '')}")
    doc.add_paragraph(f"Status: {inspection.get('status', '')}")
    doc.add_paragraph(f"Result: {inspection.get('result_status', '')}")

    doc.add_heading("Extracted Declarations", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    for key, value in declaration.items():
      row = table.add_row().cells
      row[0].text = key.replace("_", " ")
      row[1].text = clean(value) or "Not Detected"

    doc.add_heading("Compliance Checks", level=1)
    checks_table = doc.add_table(rows=1, cols=5)
    checks_table.style = "Table Grid"
    for idx, heading in enumerate(["Field", "Status", "Rule", "Source", "Confidence"]):
      checks_table.rows[0].cells[idx].text = heading
    for row_values in build_rows(detail)[1:]:
      row = checks_table.add_row().cells
      for idx, value in enumerate(row_values):
        row[idx].text = value

    doc.add_heading("Officer Verification", level=1)
    for check in detail.get("checks", []):
      if check.get("officer_status"):
        doc.add_paragraph(f"{check.get('label')}: {check.get('officer_status')} - {check.get('officer_remark') or ''}")

    doc.save(str(output_path))


def docx_pt(value):
    from docx.shared import Pt
    return Pt(value)


def main():
    if len(sys.argv) != 4:
        print("Usage: generate_report.py <input.json> <output.pdf> <output.docx>", file=sys.stderr)
        sys.exit(1)
    detail = json.loads(Path(sys.argv[1]).read_text())
    create_pdf(detail, sys.argv[2])
    create_docx(detail, sys.argv[3])
    print(json.dumps({"pdf": sys.argv[2], "docx": sys.argv[3]}))


if __name__ == "__main__":
    main()
